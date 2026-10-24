#!/usr/bin/env python3
"""Extract study material into Markdown with citation anchors.

Studying with Claude cites every fact by source and location, so this script
keeps the location of each piece of text:

  PDF          ## [p.3]              cite as [S1 p.3]
  PowerPoint   ## [slide 5] Title    cite as [S1 slide 5]
  Word, text,  ## [§2] Heading       cite as [S1 §2]
  HTML         [¶12] paragraph       cite as [S1 ¶12]
  YouTube      [00:12:34] text       cite as [S1 12:34]

Usage:
  python extract.py lecture.pdf --id S1 --out study/bio-101/sources/S1-lecture.md
  python extract.py "https://www.youtube.com/watch?v=VIDEO_ID" --id S2 --lang en,pt

Optional dependencies, installed only for the formats you use:
  pip install pypdf python-docx python-pptx youtube-transcript-api

Exit codes: 0 success, 1 error, 2 a required package is missing.
"""

from __future__ import annotations

import argparse
import datetime as dt
import html
import re
import sys
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import parse_qs, urlparse

MISSING_DEP = 2
LOW_TEXT_CHARS = 25  # a page with less text than this is probably scanned


class MissingDependency(Exception):
    def __init__(self, package: str, purpose: str):
        super().__init__(f"{purpose} needs the '{package}' package. Install it with:\n    pip install {package}")


# ---------------------------------------------------------------- helpers

def clean(text: str) -> str:
    text = text.replace(" ", " ").replace("\r", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def md_table(rows: list[list[str]]) -> str:
    rows = [[clean(c).replace("\n", " ").replace("|", "\\|") for c in r] for r in rows if any(c.strip() for c in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def timestamp(seconds: float) -> str:
    seconds = int(seconds)
    return f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"


class Doc:
    """Collects Markdown blocks and numbers paragraphs and sections."""

    def __init__(self) -> None:
        self.blocks: list[str] = []
        self.para = 0
        self.section = 0
        self.warnings: list[str] = []

    def heading(self, text: str, level: int = 2) -> None:
        text = clean(text)
        if text:
            self.section += 1
            self.blocks.append(f"{'#' * max(2, min(level, 6))} [§{self.section}] {text}")

    def paragraph(self, text: str, bullet: str = "") -> None:
        text = clean(text)
        if text:
            self.para += 1
            self.blocks.append(f"{bullet}[¶{self.para}] {text}")

    def raw(self, text: str) -> None:
        if text.strip():
            self.blocks.append(text.rstrip())


# ---------------------------------------------------------------- PDF

def extract_pdf(path: Path, doc: Doc) -> dict:
    pages: list[str] = []
    try:
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(str(path))
        pages = [(p.extract_text() or "") for p in reader.pages]
    except ImportError:
        try:
            import pdfplumber  # type: ignore

            with pdfplumber.open(str(path)) as pdf:
                pages = [(p.extract_text() or "") for p in pdf.pages]
        except ImportError:
            try:
                import fitz  # type: ignore  (PyMuPDF)

                with fitz.open(str(path)) as pdf:
                    pages = [p.get_text() for p in pdf]
            except ImportError:
                raise MissingDependency("pypdf", "Reading PDF files")

    low = []
    for number, text in enumerate(pages, start=1):
        doc.raw(f"## [p.{number}]")
        text = clean(text)
        if len(text) < LOW_TEXT_CHARS:
            low.append(number)
            doc.raw("_(Little or no extractable text on this page. It is probably scanned or an image: "
                    "read the page visually and transcribe it, marking unclear words as [illegible].)_")
        if text:
            doc.raw(text)
    if low:
        doc.warnings.append(f"Pages with little or no text (probably scanned): {', '.join(map(str, low))}")
    return {"type": "pdf", "pages": len(pages)}


# ---------------------------------------------------------------- Word

def extract_docx(path: Path, doc: Doc) -> dict:
    try:
        import docx  # type: ignore
        from docx.table import Table  # type: ignore
        from docx.text.paragraph import Paragraph  # type: ignore
    except ImportError:
        raise MissingDependency("python-docx", "Reading Word files")

    document = docx.Document(str(path))
    images = 0
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            style = (para.style.name if para.style is not None else "") or ""
            if child.xpath(".//*[local-name()='drawing' or local-name()='pict']"):
                images += 1
            if style.startswith("Heading") or style == "Title":
                digits = re.findall(r"\d+", style)
                doc.heading(para.text, 1 + (int(digits[0]) if digits else 1))
            elif "List" in style:
                doc.paragraph(para.text, bullet="- ")
            else:
                doc.paragraph(para.text)
        elif tag == "tbl":
            table = Table(child, document)
            doc.raw(md_table([[cell.text for cell in row.cells] for row in table.rows]))
    if images:
        doc.warnings.append(f"{images} embedded image(s) or drawing(s) were not extracted; ask for them if they matter.")
    return {"type": "docx", "paragraphs": doc.para}


# ---------------------------------------------------------------- PowerPoint

def extract_pptx(path: Path, doc: Doc) -> dict:
    try:
        from pptx import Presentation  # type: ignore
        from pptx.enum.shapes import MSO_SHAPE_TYPE  # type: ignore
    except ImportError:
        raise MissingDependency("python-pptx", "Reading PowerPoint files")

    presentation = Presentation(str(path))

    def shape_type(shape):
        try:
            return shape.shape_type
        except NotImplementedError:  # some shapes have no recognized type
            return None

    def walk(shapes, chunks: list[list[str]], counts: dict) -> None:
        """Collect bullet lines into chunks; each table is a chunk of its own."""
        for shape in shapes:
            kind = shape_type(shape)
            if kind == MSO_SHAPE_TYPE.GROUP:
                walk(shape.shapes, chunks, counts)
                continue
            if kind == MSO_SHAPE_TYPE.PICTURE:
                counts["images"] += 1
            if getattr(shape, "has_table", False) and shape.has_table:
                chunks.append([md_table([[cell.text for cell in row.cells] for row in shape.table.rows])])
                chunks.append([])
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                if shape == counts.get("title_shape"):
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    text = clean("".join(run.text for run in paragraph.runs))
                    if text:
                        chunks[-1].append("  " * paragraph.level + "- " + text)

    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = clean(title_shape.text) if title_shape is not None and title_shape.has_text_frame else ""
        doc.raw(f"## [slide {number}] {title}".rstrip())
        chunks: list[list[str]] = [[]]
        counts = {"images": 0, "title_shape": title_shape}
        walk(slide.shapes, chunks, counts)
        for chunk in chunks:
            if chunk:
                doc.raw("\n".join(chunk))
        if counts["images"]:
            doc.raw(f"_(This slide has {counts['images']} image(s). View the slide if a diagram matters.)_")
        if slide.has_notes_slide:
            notes = clean(slide.notes_slide.notes_text_frame.text or "")
            if notes:
                doc.raw("**Speaker notes:** " + notes)
    return {"type": "pptx", "slides": len(presentation.slides)}


# ---------------------------------------------------------------- text, Markdown, HTML

def extract_text(path: Path, doc: Doc) -> dict:
    text = path.read_text(encoding="utf-8", errors="replace")
    for block in re.split(r"\n\s*\n", text):
        block = block.strip()
        if not block:
            continue
        match = re.match(r"^(#{1,6})\s+(.*)$", block)
        if match and "\n" not in block:
            doc.heading(match.group(2), len(match.group(1)) + 1)
        else:
            doc.paragraph(block)
    return {"type": "text", "paragraphs": doc.para}


class _HTMLText(HTMLParser):
    SKIP = {"script", "style", "nav", "footer", "header", "aside", "noscript", "svg", "form"}
    BLOCKS = {"p", "li", "blockquote", "pre", "td", "th", "dd", "dt", "figcaption"}

    def __init__(self, doc: Doc) -> None:
        super().__init__(convert_charrefs=True)
        self.doc = doc
        self.skip_depth = 0
        self.buffer: list[str] = []
        self.current: str | None = None
        self.title = ""
        self.in_title = False

    def flush(self) -> None:
        text = clean(" ".join(self.buffer))
        self.buffer = []
        if not text or self.current is None:
            return
        if self.current.startswith("h") and self.current[1:].isdigit():
            self.doc.heading(text, int(self.current[1:]) + 1)
        else:
            self.doc.paragraph(text, bullet="- " if self.current == "li" else "")

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip_depth += 1
        elif tag == "title":
            self.in_title = True
        elif tag in self.BLOCKS or re.fullmatch(r"h[1-6]", tag):
            self.flush()
            self.current = tag

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self.skip_depth = max(0, self.skip_depth - 1)
        elif tag == "title":
            self.in_title = False
        elif tag in self.BLOCKS or re.fullmatch(r"h[1-6]", tag):
            self.flush()
            self.current = None

    def handle_data(self, data):
        if self.in_title:
            self.title += data
        elif not self.skip_depth and self.current is not None:
            self.buffer.append(data)


def extract_html(path: Path, doc: Doc) -> dict:
    parser = _HTMLText(doc)
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    parser.flush()
    return {"type": "web", "paragraphs": doc.para, "page_title": html.unescape(clean(parser.title))}


# ---------------------------------------------------------------- YouTube

def youtube_id(url: str) -> str | None:
    parsed = urlparse(url)
    host = (parsed.hostname or "").lower()
    if host.endswith("youtu.be"):
        return parsed.path.strip("/").split("/")[0] or None
    if "youtube" in host:
        if parsed.path == "/watch":
            return parse_qs(parsed.query).get("v", [None])[0]
        match = re.match(r"^/(shorts|embed|live|v)/([^/?#]+)", parsed.path)
        if match:
            return match.group(2)
    return None


def extract_youtube(url: str, doc: Doc, languages: list[str], group_seconds: float = 30.0) -> dict:
    video = youtube_id(url)
    if not video:
        raise ValueError(f"Not a recognizable YouTube URL: {url}")
    try:
        from youtube_transcript_api import YouTubeTranscriptApi  # type: ignore
    except ImportError:
        raise MissingDependency("youtube-transcript-api", "Fetching YouTube transcripts")

    if hasattr(YouTubeTranscriptApi, "get_transcript"):  # versions before 1.0
        snippets = [(s["start"], s["text"]) for s in YouTubeTranscriptApi.get_transcript(video, languages=languages)]
    else:
        fetched = YouTubeTranscriptApi().fetch(video, languages=languages)
        snippets = [(s.start, s.text) for s in fetched]

    start, parts = None, []
    for begin, text in snippets:
        text = clean(text.replace("\n", " "))
        if not text:
            continue
        if start is None:
            start = begin
        parts.append(text)
        if begin - start >= group_seconds:
            doc.raw(f"[{timestamp(start)}] " + " ".join(parts))
            start, parts = None, []
    if parts:
        doc.raw(f"[{timestamp(start or 0)}] " + " ".join(parts))
    doc.warnings.append("YouTube captions can be auto-generated and contain errors; check key terms against the video.")
    duration = timestamp(snippets[-1][0]) if snippets else "00:00:00"
    return {"type": "youtube", "video_id": video, "duration": duration}


# ---------------------------------------------------------------- main

EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".pptx": extract_pptx,
    ".txt": extract_text,
    ".md": extract_text,
    ".markdown": extract_text,
    ".html": extract_html,
    ".htm": extract_html,
}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Extract study material into Markdown with citation anchors.")
    parser.add_argument("input", help="File path (.pdf .docx .pptx .txt .md .html) or a YouTube URL")
    parser.add_argument("--id", default="S?", help="Source ID to put in the header, e.g. S3")
    parser.add_argument("--title", help="Human-readable title (defaults to the file name)")
    parser.add_argument("--out", help="Write Markdown here instead of printing it")
    parser.add_argument("--lang", default="en", help="YouTube caption languages in order of preference, e.g. en,pt")
    args = parser.parse_args(argv)

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")

    doc = Doc()
    is_url = re.match(r"^https?://", args.input) is not None
    try:
        if is_url:
            if not youtube_id(args.input):
                print("Only YouTube URLs are handled here. Fetch other web pages with your web tool, "
                      "or save the page as .html and pass the file.", file=sys.stderr)
                return 1
            meta = extract_youtube(args.input, doc, [l.strip() for l in args.lang.split(",") if l.strip()])
            title = args.title or f"YouTube video {meta['video_id']}"
        else:
            path = Path(args.input)
            if not path.is_file():
                print(f"File not found: {path}", file=sys.stderr)
                return 1
            extractor = EXTRACTORS.get(path.suffix.lower())
            if extractor is None:
                print(f"Unsupported file type '{path.suffix}'. Supported: {', '.join(sorted(EXTRACTORS))} "
                      "and YouTube URLs. Images: read them directly instead.", file=sys.stderr)
                return 1
            meta = extractor(path, doc)
            title = args.title or meta.pop("page_title", None) or path.stem.replace("_", " ").replace("-", " ")
    except MissingDependency as exc:
        print(str(exc), file=sys.stderr)
        return MISSING_DEP
    except Exception as exc:  # report a readable error instead of a traceback
        print(f"Could not extract {args.input}: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1

    meta.pop("page_title", None)
    header = ["---", f"source_id: {args.id}", f"title: {title}", f"original: {args.input}"]
    header += [f"{key}: {value}" for key, value in meta.items()]
    header += [f"extracted: {dt.date.today().isoformat()}", "---", "", f"# {args.id} · {title}", ""]
    body = "\n\n".join(doc.blocks) if doc.blocks else "_(No text could be extracted.)_"
    output = "\n".join(header) + "\n" + body + "\n"

    if args.out:
        out = Path(args.out)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(output, encoding="utf-8")
        print(f"Wrote {out} ({len(doc.blocks)} blocks)", file=sys.stderr)
    else:
        sys.stdout.write(output)
    for warning in doc.warnings:
        print("Note: " + warning, file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
